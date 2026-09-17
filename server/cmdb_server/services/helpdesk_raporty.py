"""Raporty czasu pracy helpdesku i ich eksport.

Dwa raporty sa jednym zestawieniem ogladanym z dwoch stron: firma -> technik
-> zgloszenie i technik -> firma -> zgloszenie. Dlatego licza je dwie funkcje
o tym samym ksztalcie wyniku, a eksport i szablony traktuja je tak samo.

Minuty sumujemy z tabeli czasu, nie z czegokolwiek wyliczanego z watku:
czas pracy jest rejestrem zdarzen, wiec raport ma byc jego prostym podsumowaniem
i ma sie zgadzac z tym, co technik wpisal.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..models import STATUSY_ZGLOSZENIA, CzasPracy, PortalUser, Tenant, Zgloszenie
from . import helpdesk


@dataclass
class PozycjaZgloszenia:
    numer: str
    temat: str
    status: str
    minuty: int


@dataclass
class Grupa:
    """Technik w raporcie firmy albo firma w raporcie technika."""

    id: str
    nazwa: str
    minuty: int
    zgloszenia: list[PozycjaZgloszenia] = field(default_factory=list)


@dataclass
class Raport:
    tytul: str
    podtytul: str
    suma: int
    grupy: list[Grupa] = field(default_factory=list)

    @property
    def zgloszen(self) -> int:
        """Ile roznych zgloszen sklada sie na ten raport.

        Jedno zgloszenie bywa u dwoch technikow - liczymy je raz, bo inaczej
        "9 zgloszen" przy trzech sprawach wygladaloby na blad danych.
        """
        return len({pozycja.numer for grupa in self.grupy for pozycja in grupa.zgloszenia})


# --- okresy -----------------------------------------------------------------

# Nazwy dla oka - do nazwy pliku i tak idzie wersja bez ogonkow z nazwa_pliku().
MIESIACE = (
    "styczeń", "luty", "marzec", "kwiecień", "maj", "czerwiec",
    "lipiec", "sierpień", "wrzesień", "październik", "listopad", "grudzień",
)


def miesiac(wartosc: str | None) -> tuple[datetime, datetime, str]:
    """Zakres jednego miesiaca z napisu "2026-09". Domyslnie miesiac biezacy.

    Zwracamy przedzial domkniety z lewej i otwarty z prawej, bo tylko taki
    nie gubi ostatniej sekundy miesiaca ani nie liczy pierwszej nastepnego.
    """
    dzis = datetime.now(timezone.utc).date()
    rok, numer = dzis.year, dzis.month
    if wartosc and re.fullmatch(r"\d{4}-\d{2}", wartosc.strip()):
        rok, numer = int(wartosc[:4]), int(wartosc[5:7])
        numer = min(max(numer, 1), 12)

    poczatek = datetime(rok, numer, 1, tzinfo=timezone.utc)
    koniec = datetime(rok + (numer == 12), (numer % 12) + 1, 1, tzinfo=timezone.utc)
    return poczatek, koniec, f"{MIESIACE[numer - 1]} {rok}"


def ostatnie_miesiace(ile: int = 12) -> list[tuple[str, str]]:
    """Lista miesiecy do wyboru w formularzu: [("2026-09", "wrzesień 2026"), ...]."""
    dzis = date.today().replace(day=1)
    wybor = []
    for _ in range(ile):
        wybor.append((dzis.strftime("%Y-%m"), f"{MIESIACE[dzis.month - 1]} {dzis.year}"))
        dzis = (dzis - timedelta(days=1)).replace(day=1)
    return wybor


# --- liczenie ---------------------------------------------------------------

def _pozycje(db: Session, warunki, grupa_po):
    """Minuty w rozbiciu na grupe i zgloszenie - jedno zapytanie na raport."""
    return db.execute(
        select(
            grupa_po,
            Zgloszenie.numer_pelny,
            Zgloszenie.temat,
            Zgloszenie.status,
            func.sum(CzasPracy.minuty),
        )
        .join(Zgloszenie, Zgloszenie.id == CzasPracy.zgloszenie_id)
        .where(*warunki)
        .group_by(grupa_po, Zgloszenie.numer_pelny, Zgloszenie.temat, Zgloszenie.status)
        .order_by(Zgloszenie.numer_pelny)
    ).all()


def _zloz(wiersze, nazwy: dict[str, str]) -> list[Grupa]:
    grupy: dict[str, Grupa] = {}
    for klucz, numer, temat, status, minuty in wiersze:
        grupa = grupy.setdefault(
            klucz, Grupa(id=klucz, nazwa=nazwy.get(klucz, klucz), minuty=0)
        )
        grupa.minuty += int(minuty)
        grupa.zgloszenia.append(PozycjaZgloszenia(
            numer=numer, temat=temat,
            status=STATUSY_ZGLOSZENIA.get(status, status), minuty=int(minuty),
        ))
    # Najwiekszy wklad na gorze: raport czyta sie po to, zeby zobaczyc, gdzie
    # poszedl czas, a nie zeby przegladac alfabet.
    return sorted(grupy.values(), key=lambda g: g.minuty, reverse=True)


def raport_firmy(db: Session, tenant_id: str, okres: str | None = None) -> Raport:
    """Firma -> technik -> zgloszenie."""
    od, do, opis = miesiac(okres)
    firma = db.get(Tenant, tenant_id)

    wiersze = _pozycje(
        db,
        (CzasPracy.tenant_id == tenant_id, CzasPracy.utworzono >= od, CzasPracy.utworzono < do),
        CzasPracy.technik_id,
    )
    nazwy = {
        identyfikator: nazwa or email
        for identyfikator, nazwa, email in db.execute(
            select(PortalUser.id, PortalUser.full_name, PortalUser.email)
        ).all()
    }
    grupy = _zloz(wiersze, nazwy)
    return Raport(
        tytul=firma.name if firma else tenant_id,
        podtytul=opis,
        suma=sum(g.minuty for g in grupy),
        grupy=grupy,
    )


def raport_technika(
    db: Session, technik_id: str, okres: str | None = None,
    firmy: list[str] | None = None,
) -> Raport:
    """Technik -> firma -> zgloszenie.

    ``firmy`` zawezaja raport do wskazanych firm. Sluzy to izolacji: technik
    oglada wlasny czas, ale tylko w firmach, ktore obsluguje - inaczej widzialby
    nazwy firm, do ktorych nie ma dostepu.
    """
    od, do, opis = miesiac(okres)
    technik = db.get(PortalUser, technik_id)

    warunki = [
        CzasPracy.technik_id == technik_id,
        CzasPracy.utworzono >= od,
        CzasPracy.utworzono < do,
    ]
    if firmy is not None:
        warunki.append(CzasPracy.tenant_id.in_(firmy or ["-"]))

    wiersze = _pozycje(db, warunki, CzasPracy.tenant_id)
    nazwy = {
        identyfikator: nazwa
        for identyfikator, nazwa in db.execute(select(Tenant.id, Tenant.name)).all()
    }
    grupy = _zloz(wiersze, nazwy)
    return Raport(
        tytul=(technik.full_name or technik.email) if technik else technik_id,
        podtytul=opis,
        suma=sum(g.minuty for g in grupy),
        grupy=grupy,
    )


@dataclass
class WierszPodsumowania:
    """Jeden technik w podsumowaniu miesiaca: suma i rozbicie na firmy."""

    technik_id: str
    nazwa: str
    minuty: int
    firmy: list[tuple[str, int]]


@dataclass
class Podsumowanie:
    """Caly miesiac na jednym ekranie: technicy w wierszach, firmy w kolumnach."""

    okres: str
    nazwy_firm: list[tuple[str, str]]
    technicy: list[WierszPodsumowania]
    sumy_firm: dict[str, int]
    suma: int


def podsumowanie(
    db: Session, okres: str | None = None, firmy: list[str] | None = None,
    technicy: list[str] | None = None,
) -> Podsumowanie:
    """Czas wszystkich technikow za miesiac, bez klikania po jednym.

    Dotychczas odpowiedz na "ile kto zrobil w tym miesiacu" wymagala otwarcia
    raportu osobno dla kazdego technika i zsumowania w glowie. Tu wszystko
    jest w jednej tabeli: wiersz to technik, kolumna to firma, a ostatnia
    kolumna - jego suma miesiaca.

    ``firmy`` i ``technicy`` zawezaja zestawienie; sluzy to izolacji - technik
    oglada wlasny wiersz i tylko firmy, ktore obsluguje.
    """
    od, do, opis = miesiac(okres)

    warunki = [CzasPracy.utworzono >= od, CzasPracy.utworzono < do]
    if firmy is not None:
        warunki.append(CzasPracy.tenant_id.in_(firmy or ["-"]))
    if technicy is not None:
        warunki.append(CzasPracy.technik_id.in_(technicy or ["-"]))

    wiersze = db.execute(
        select(
            CzasPracy.technik_id,
            func.coalesce(PortalUser.full_name, PortalUser.email),
            CzasPracy.tenant_id,
            func.sum(CzasPracy.minuty),
        )
        .join(PortalUser, PortalUser.id == CzasPracy.technik_id)
        .where(*warunki)
        .group_by(CzasPracy.technik_id, PortalUser.full_name, PortalUser.email,
                  CzasPracy.tenant_id)
    ).all()

    nazwy = dict(db.execute(select(Tenant.id, Tenant.name)).all())
    zebrane: dict[str, WierszPodsumowania] = {}
    sumy_firm: dict[str, int] = {}
    uzyte_firmy: set[str] = set()

    for technik_id, nazwa, tenant_id, minuty in wiersze:
        minuty = int(minuty)
        wiersz = zebrane.setdefault(
            technik_id, WierszPodsumowania(technik_id, nazwa, 0, [])
        )
        wiersz.minuty += minuty
        wiersz.firmy.append((tenant_id, minuty))
        sumy_firm[tenant_id] = sumy_firm.get(tenant_id, 0) + minuty
        uzyte_firmy.add(tenant_id)

    # Firmy w kolumnach tylko te, w ktorych ktos w tym miesiacu pracowal -
    # pusta kolumna niczego nie mowi, a rozpycha tabele.
    kolumny = sorted(
        ((identyfikator, nazwy.get(identyfikator, identyfikator))
         for identyfikator in uzyte_firmy),
        key=lambda para: para[1].lower(),
    )
    lista = sorted(zebrane.values(), key=lambda w: (-w.minuty, w.nazwa.lower()))
    return Podsumowanie(
        okres=opis,
        nazwy_firm=kolumny,
        technicy=lista,
        sumy_firm=sumy_firm,
        suma=sum(w.minuty for w in lista),
    )


def podsumowanie_do_raportu(dane: Podsumowanie) -> Raport:
    """Podsumowanie w postaci, ktora rozumie eksport CSV/XLSX."""
    grupy = [
        Grupa(
            id=wiersz.technik_id, nazwa=wiersz.nazwa, minuty=wiersz.minuty,
            zgloszenia=[
                PozycjaZgloszenia(
                    numer="", temat=dict(dane.nazwy_firm).get(tenant_id, tenant_id),
                    status="", minuty=minuty,
                )
                for tenant_id, minuty in sorted(wiersz.firmy, key=lambda para: -para[1])
            ],
        )
        for wiersz in dane.technicy
    ]
    return Raport(
        tytul="Czas techników", podtytul=dane.okres,
        suma=dane.suma, grupy=grupy,
    )


# --- eksport ----------------------------------------------------------------

NAGLOWKI = ("Grupa", "Zgłoszenie", "Temat", "Status", "Minuty", "Czas")


def wiersze_eksportu(raport: Raport) -> list[tuple]:
    """Raport splaszczony do tabeli.

    Minuty ida obok czasu opisowego: pierwsze da sie zsumowac w arkuszu,
    drugie da sie przeczytac. Sama kolumna "1 h 35 min" zamienia arkusz
    w obrazek.
    """
    wiersze = []
    for grupa in raport.grupy:
        for pozycja in grupa.zgloszenia:
            wiersze.append((
                grupa.nazwa, pozycja.numer, pozycja.temat, pozycja.status,
                pozycja.minuty, helpdesk.formatuj_czas(pozycja.minuty),
            ))
        wiersze.append((
            grupa.nazwa, "", "RAZEM", "", grupa.minuty, helpdesk.formatuj_czas(grupa.minuty),
        ))
    wiersze.append((
        "RAZEM", "", raport.podtytul, "", raport.suma, helpdesk.formatuj_czas(raport.suma),
    ))
    return wiersze


def do_csv(raport: Raport) -> bytes:
    """CSV dla arkuszy kalkulacyjnych w polskich ustawieniach.

    Srednik zamiast przecinka i BOM na poczatku: Excel w polskiej wersji
    inaczej wrzuca caly wiersz do jednej kolumny i psuje ogonki.
    """
    bufor = io.StringIO()
    zapis = csv.writer(bufor, delimiter=";", lineterminator="\r\n")
    zapis.writerow(NAGLOWKI)
    zapis.writerows(wiersze_eksportu(raport))
    return bufor.getvalue().encode("utf-8-sig")


def do_xlsx(raport: Raport) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    skoroszyt = Workbook()
    arkusz = skoroszyt.active
    arkusz.title = "Czas pracy"

    arkusz.append([f"{raport.tytul} - {raport.podtytul}"])
    arkusz["A1"].font = Font(bold=True, size=13)
    arkusz.append([])
    arkusz.append(list(NAGLOWKI))
    for komorka in arkusz[3]:
        komorka.font = Font(bold=True)

    for wiersz in wiersze_eksportu(raport):
        arkusz.append(list(wiersz))
        if wiersz[2] in ("RAZEM", raport.podtytul) and not wiersz[1]:
            for komorka in arkusz[arkusz.max_row]:
                komorka.font = Font(bold=True)

    for kolumna, szerokosc in zip("ABCDEF", (28, 14, 56, 14, 10, 14)):
        arkusz.column_dimensions[kolumna].width = szerokosc
    for wiersz in arkusz.iter_rows(min_row=4, min_col=5, max_col=5):
        for komorka in wiersz:
            komorka.alignment = Alignment(horizontal="right")

    bufor = io.BytesIO()
    skoroszyt.save(bufor)
    return bufor.getvalue()


def nazwa_pliku(raport: Raport, rozszerzenie: str) -> str:
    """Nazwa pliku bez ogonkow i spacji - przechodzi przez kazdy system."""
    podstawa = f"czas-{raport.tytul}-{raport.podtytul}".lower()
    podstawa = (podstawa.replace("ł", "l").replace("ą", "a").replace("ę", "e")
                .replace("ś", "s").replace("ć", "c").replace("ż", "z")
                .replace("ź", "z").replace("ó", "o").replace("ń", "n"))
    podstawa = re.sub(r"[^a-z0-9]+", "-", podstawa).strip("-")
    return f"{podstawa}.{rozszerzenie}"
